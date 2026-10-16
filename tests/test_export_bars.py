"""Exporting every ES bar to three spreadsheets.

Four things would make the sheets wrong in ways nobody notices until a
study built on them gives a strange answer.

A BAR MUST NOT SPAN A CONTRACT ROLL. Across one the price steps by carry,
so a bar straddling it carries a high and a low from two different
instruments and shows a range nobody traded.

AGGREGATION MUST KEEP THE EXTREMES. Open of the first, high and low of
everything inside, close of the last. Taking the latest high instead of
the running maximum quietly throws away the bar's range.

THE TWO SOURCES MUST NOT DOUBLE UP. The OHLCV file and the depth files
overlap by days. Two aggregations of the same second disagree in the
fourth decimal, and merging them puts both in the sheet.

AND THE CLOCK GRID IS THE CLOCK. Bars starting wherever the data happens
to begin put every timeframe on its own private grid.
"""

from __future__ import annotations

import csv
import importlib.util
import pathlib
import sys

import pytest

ROOT = pathlib.Path(__file__).resolve().parent.parent


def _load():
    path = ROOT / "scripts" / "export_bars.py"
    spec = importlib.util.spec_from_file_location("export_bars", path)
    mod = importlib.util.module_from_spec(spec)
    sys.modules["export_bars"] = mod
    spec.loader.exec_module(mod)
    return mod


eb = _load()
NS = 1_000_000_000
PX = 1_000_000_000
T0 = 1_727_654_400.0            # 2024-09-30 00:00 UTC


def rec(ts, o, h, l, c, v=10, iid=1):
    return {"ts_event": int(ts * NS), "open": int(o * PX), "high": int(h * PX),
            "low": int(l * PX), "close": int(c * PX), "volume": v,
            "instrument_id": iid}


def trade_rec(ts, px, sz=2, iid=9):
    return {"ts_event": int(ts * NS), "action": "T", "side": "B",
            "price": int(px * PX), "size": sz, "instrument_id": iid}


# --------------------------------------------------------- the builder

def test_bars_sit_on_the_clock_not_on_the_first_print():
    b = eb.Builder(60)
    b.feed(T0 + 70, 10, 12, 9, 11, 5, 1)
    b.feed(T0 + 71, 11, 13, 10, 12, 5, 1)
    b.feed(T0 + 130, 20, 21, 19, 20, 5, 1)
    b.finish()
    assert len(b.rows) == 2
    # the first bar opened in the 60-second slot, not at 70
    assert b.rows[0][1].endswith(":00")


def test_aggregation_keeps_the_extremes_and_the_last_close():
    """The extremes must come from ANY bar inside, not the latest one.

    Ordering the fixture so the last bar happens to hold both the high
    and the low lets `cur.high = h` pass as if it were `max(...)`, and
    the sheet quietly loses every bar's real range.
    """
    b = eb.Builder(60)
    b.feed(T0, 10, 15, 5, 12, 3, 1)        # the extremes are both HERE
    b.feed(T0 + 1, 12, 13, 8, 9, 4, 1)     # and this bar sits inside them
    b.finish()
    d, t, price, o, h, lo, vol, contract, u = b.rows[0]
    assert o == 10, "open is not the first bar's"
    assert h == 15, "high was taken from the last bar, not the extreme"
    assert lo == 5, "low was taken from the last bar, not the extreme"
    assert price == 9, "close is not the last bar's"
    assert vol == 7


def test_the_extremes_survive_whichever_bar_holds_them():
    """The mirror of the above: now the LAST bar has both."""
    b = eb.Builder(60)
    b.feed(T0, 10, 12, 9, 11, 1, 1)
    b.feed(T0 + 1, 11, 20, 2, 15, 1, 1)
    b.finish()
    _, _, price, o, h, lo, _, _, _ = b.rows[0]
    assert (o, h, lo, price) == (10, 20, 2, 15)


def test_price_is_the_close_not_the_open_or_the_midpoint():
    b = eb.Builder(60)
    b.feed(T0, 100, 110, 90, 105, 1, 1)
    b.finish()
    assert b.rows[0][2] == pytest.approx(105)


def test_a_change_of_contract_closes_the_bar_mid_slot():
    """THE roll guard. Both prints fall in the same minute, but they are
    different instruments and a bar across them is carry, not a range."""
    b = eb.Builder(60)
    b.feed(T0, 100, 100, 100, 100, 1, 111)
    b.feed(T0 + 1, 130, 130, 130, 130, 1, 222)
    b.finish()
    assert len(b.rows) == 2
    assert b.rows[0][7] == 111 and b.rows[1][7] == 222
    assert b.rows[0][4] == pytest.approx(100)     # high is not 130


def test_a_bar_within_one_contract_still_merges():
    b = eb.Builder(60)
    b.feed(T0, 100, 101, 99, 100, 1, 111)
    b.feed(T0 + 1, 100, 105, 95, 102, 1, 111)
    b.finish()
    assert len(b.rows) == 1
    assert b.rows[0][4] == pytest.approx(105)
    assert b.rows[0][5] == pytest.approx(95)


def test_volume_sums_across_the_bar():
    b = eb.Builder(300)
    for i in range(5):
        b.feed(T0 + i, 100, 100, 100, 100, 7, 1)
    b.finish()
    assert b.rows[0][6] == 35


def test_finishing_twice_does_not_duplicate_the_last_bar():
    b = eb.Builder(60)
    b.feed(T0, 100, 100, 100, 100, 1, 1)
    b.finish()
    b.finish()
    assert len(b.rows) == 1


def test_the_three_timeframes_are_the_ones_asked_for():
    assert [n for n, _ in eb.FRAMES] == ["1m", "5m", "15m"]
    assert [s for _, s in eb.FRAMES] == [60, 300, 900]


def test_a_coarser_frame_holds_fewer_bars_over_the_same_prints():
    one, five = eb.Builder(60), eb.Builder(300)
    for i in range(600):
        for b in (one, five):
            b.feed(T0 + i, 100, 100, 100, 100, 1, 1)
    one.finish(); five.finish()
    assert len(one.rows) == 10
    assert len(five.rows) == 2


# ------------------------------------------------------- the columns

def test_the_requested_columns_come_first_and_in_order():
    """date, price and time were what was asked for; the rest follow."""
    assert eb.HEADERS[:3] == ("date", "time", "price")
    assert set(eb.HEADERS) >= {"open", "high", "low", "volume",
                               "contract", "utc"}


def test_date_and_time_are_exchange_time_and_utc_is_separate():
    d, t, u = eb.stamp(T0)
    assert u.endswith("Z")
    assert u.startswith("2024-09-30T00:00")
    assert d in ("2024-09-29", "2024-09-30")     # Chicago is behind UTC
    assert len(t.split(":")) == 3


def test_the_time_column_is_when_the_bar_opened():
    b = eb.Builder(900)
    b.feed(T0 + 800, 1, 1, 1, 1, 1, 1)           # lands in the 0-900 slot
    b.finish()
    assert b.rows[0][1].endswith(":00:00") or b.rows[0][1].endswith(":00")


# -------------------------------------------------- the two sources

def test_the_ohlcv_file_fills_the_builders():
    b = eb.Builder(60)
    s = eb.Source(path="x", schema="ohlcv-1s")
    import liqmap.dbn as dbn
    orig = dbn.open_file
    dbn.open_file = lambda p: [rec(T0 + i, 100, 101, 99, 100)
                               for i in range(120)]
    try:
        eb.read_ohlcv("x", [b], s)
    finally:
        dbn.open_file = orig
    b.finish()
    assert s.kept == 120 and s.skipped == 0
    assert len(b.rows) == 2


def _read(records):
    b = eb.Builder(60)
    s = eb.Source(path="x")
    import liqmap.dbn as dbn
    orig = dbn.open_file
    dbn.open_file = lambda p: records
    try:
        eb.read_ohlcv("x", [b], s)
    finally:
        dbn.open_file = orig
    b.finish()
    return b, s


def test_malformed_ohlcv_rows_are_skipped_not_repaired():
    b, s = _read([rec(T0, 100, 99, 101, 100),       # high below low
                  rec(T0 + 1, 100, 101, 99, 100),   # fine
                  rec(T0 + 2, 105, 101, 99, 100)])  # open above the high
    assert s.kept == 1 and s.skipped == 2


def test_a_row_with_no_price_is_skipped():
    """All zeros passes every range test -- 0 <= 0 <= 0 -- so only the
    price guard catches it, and a zero bar in the sheet is a 5,000-point
    move on the row either side of it."""
    b, s = _read([rec(T0, 0, 0, 0, 0), rec(T0 + 1, 100, 101, 99, 100)])
    assert s.kept == 1 and s.skipped == 1
    assert all(r[2] > 0 for r in b.rows)


def test_a_row_with_an_unreadable_stamp_is_skipped():
    bad = {"ts_event": None, "open": int(100 * PX), "high": int(101 * PX),
           "low": int(99 * PX), "close": int(100 * PX), "volume": 1,
           "instrument_id": 1}
    b, s = _read([bad, rec(T0, 100, 101, 99, 100)])
    assert s.kept == 1 and s.skipped == 1


def test_a_close_outside_its_own_range_is_skipped():
    """The open can sit inside the range while the close does not, so
    checking only the open lets a corrupt bar through."""
    b, s = _read([rec(T0, 100, 101, 99, 140),       # close above the high
                  rec(T0 + 1, 100, 101, 99, 50),    # close below the low
                  rec(T0 + 2, 100, 101, 99, 100)])  # fine
    assert s.kept == 1 and s.skipped == 2


def test_depth_trades_inside_the_covered_span_are_skipped():
    """The OHLCV file is the authority. Merging a second aggregation of
    the same second puts two slightly different bars in the sheet."""
    b = eb.Builder(60)
    s = eb.Source(path="d", schema="mbp-10")
    import liqmap.dbn as dbn
    orig = dbn.open_file
    dbn.open_file = lambda p: [trade_rec(T0 + 10, 5000.0),
                               trade_rec(T0 + 20, 5000.5)]
    try:
        eb.read_trades("d", [b], s, covered=(T0, T0 + 100))
    finally:
        dbn.open_file = orig
    assert s.kept == 0 and s.skipped == 2
    assert s.added_days == set()


def test_depth_trades_outside_the_covered_span_are_kept():
    b = eb.Builder(60)
    s = eb.Source(path="d", schema="mbp-10")
    import liqmap.dbn as dbn
    orig = dbn.open_file
    later = T0 + 10_000_000
    dbn.open_file = lambda p: [trade_rec(later, 5000.0),
                               trade_rec(later + 1, 5000.5)]
    try:
        eb.read_trades("d", [b], s, covered=(T0, T0 + 100))
    finally:
        dbn.open_file = orig
    b.finish()
    assert s.kept == 2
    assert len(s.added_days) == 1
    assert len(b.rows) == 1
    assert b.rows[0][2] == pytest.approx(5000.5)      # close of the two


def test_with_no_ohlcv_coverage_every_trade_is_kept():
    b = eb.Builder(60)
    s = eb.Source(path="d", schema="mbp-10")
    import liqmap.dbn as dbn
    orig = dbn.open_file
    dbn.open_file = lambda p: [trade_rec(T0, 5000.0)]
    try:
        eb.read_trades("d", [b], s, covered=(None, 0.0))
    finally:
        dbn.open_file = orig
    assert s.kept == 1


def test_a_non_trade_record_is_not_counted_as_skipped_data():
    """A book update is not a malformed trade; counting it as skipped
    would report millions of 'problems' in a healthy file."""
    b = eb.Builder(60)
    s = eb.Source(path="d", schema="mbp-10")
    import liqmap.dbn as dbn
    orig = dbn.open_file
    dbn.open_file = lambda p: [{"ts_event": int(T0 * NS), "action": "A",
                                "side": "B", "price": int(5000 * PX),
                                "size": 1, "instrument_id": 9}]
    try:
        eb.read_trades("d", [b], s, covered=(None, 0.0))
    finally:
        dbn.open_file = orig
    assert s.kept == 0 and s.skipped == 0


# ------------------------------------------------------------ writing

def test_the_csv_has_a_header_and_every_row():
    rows = [("2024-09-30", "08:30:00", 5000.0, 5000.0, 5001.0, 4999.0,
             10, 1, "2024-09-30T13:30:00Z")] * 3
    import tempfile
    with tempfile.TemporaryDirectory() as d:
        p = pathlib.Path(d) / "x.csv"
        assert eb.write_csv(p, rows) == 3
        got = list(csv.reader(open(p)))
    assert got[0] == list(eb.HEADERS)
    assert len(got) == 4


def test_a_sheet_past_excels_row_limit_falls_back_to_csv():
    """Excel stops at 1,048,576 rows. Writing past it produces a file
    that will not open, which is worse than not writing it."""
    import tempfile

    class Fake(list):
        def __len__(self):
            return eb.XLSX_ROW_LIMIT + 5

    with tempfile.TemporaryDirectory() as d:
        note = eb.write_xlsx(pathlib.Path(d) / "x.xlsx", Fake(), "s")
    assert "CSV only" in note
    assert "limit" in note


def test_the_row_limit_is_excels_real_one():
    assert eb.XLSX_ROW_LIMIT == 1_048_576


def test_a_missing_openpyxl_is_reported_rather_than_crashing(monkeypatch):
    import builtins
    real = builtins.__import__

    def no_openpyxl(name, *a, **k):
        if name.startswith("openpyxl"):
            raise ImportError("nope")
        return real(name, *a, **k)

    monkeypatch.setattr(builtins, "__import__", no_openpyxl)
    note = eb.write_xlsx(pathlib.Path("/tmp/none.xlsx"), [], "s")
    assert "openpyxl" in note and "CSV only" in note


@pytest.mark.skipif(importlib.util.find_spec("openpyxl") is None,
                    reason="openpyxl not installed here")
def test_the_xlsx_round_trips_its_values():
    import tempfile
    from openpyxl import load_workbook
    rows = [("2024-09-30", "08:30:00", 5000.25, 5000.0, 5001.0, 4999.0,
             10, 1, "2024-09-30T13:30:00Z")]
    with tempfile.TemporaryDirectory() as d:
        p = pathlib.Path(d) / "x.xlsx"
        eb.write_xlsx(p, rows, "ES 1m")
        wb = load_workbook(p)
        ws = wb.active
        assert [c.value for c in ws[1]] == list(eb.HEADERS)
        assert ws["C2"].value == pytest.approx(5000.25)
        assert ws.freeze_panes == "A2"


# ------------------------------------------------- no strategy, no orders

def test_it_computes_no_signal_and_scores_no_trade():
    import inspect
    src = inspect.getsource(eb).lower()
    for word in ("win_rate", "breakeven", "first_touch", "signal",
                 "target", "stop_ticks"):
        assert word not in src, word


def test_nothing_here_places_an_order():
    import inspect
    src = inspect.getsource(eb).lower()
    for word in ("def submit", "def execute", "def place", "def order"):
        assert word not in src


def test_no_production_module_is_mutated():
    import ast
    import inspect
    tree = ast.parse(inspect.getsource(eb))
    for node in ast.walk(tree):
        if isinstance(node, ast.Assign):
            for t in node.targets:
                if isinstance(t, ast.Attribute):
                    assert getattr(t.value, "id", "") != "dbn"


def test_progress_never_writes_to_the_report():
    import io
    buf = io.StringIO()
    p = eb.Progress("x", every=2, out=buf)
    p.tick(); p.tick(); p.done()
    assert buf.getvalue()
    assert eb.Progress("x").out is sys.stderr


def test_the_row_limit_is_reported_even_without_openpyxl(monkeypatch):
    """The two reasons for "CSV only" are not interchangeable.

    On a machine without openpyxl the import guard used to answer
    first, so an oversized sheet was reported as a missing library --
    sending someone to `pip install` a package that would not have
    helped. The unfixable reason has to win.
    """
    import builtins
    import tempfile
    real = builtins.__import__

    def no_openpyxl(name, *a, **k):
        if name.startswith("openpyxl"):
            raise ImportError("not installed")
        return real(name, *a, **k)

    class Fake(list):
        def __len__(self):
            return eb.XLSX_ROW_LIMIT + 5

    monkeypatch.setattr(builtins, "__import__", no_openpyxl)
    with tempfile.TemporaryDirectory() as d:
        note = eb.write_xlsx(pathlib.Path(d) / "x.xlsx", Fake(), "s")
    assert "limit" in note
    assert "openpyxl" not in note
